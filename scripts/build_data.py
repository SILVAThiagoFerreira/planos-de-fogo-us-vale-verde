#!/usr/bin/env python3
"""Build the dashboard's small, public data files from the source workbooks."""

from __future__ import annotations

import argparse
import json
import math
import re
import unicodedata
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ASPECTS = (
    ("Diâmetro", ("Diâmetro",), ("Diâmetro (pol)", "Diâmetro")),
    ("Afastamento", ("Afastamento",), ("Afastamento (m)", "Afastamento")),
    ("Espaçamento", ("Espaçamento",), ("Espaçamento (m)", "Espaçamento")),
    ("Sub-furação", ("Sub.furação", "Subfuração", "Sub-Furação"), ("Sub-Furação (m)", "Sub-Furação", "Subfuração")),
    ("Tampão", ("Tampão",), ("Tampão Utilizado (m)", "Tampão")),
    ("Média de Perfuração", ("Média de Perfuração",), ("Média de Perfuração (m)", "Média de Perfuração")),
    ("Perfuração Específica", ("Perfuração Específica",), ("Perfuração Específica (m/m³)", "Perfuração Específica")),
    ("Nº de Furos", ("Nº de Furos", "Numero de Furos"), ("Nº de Furos", "Numero de Furos")),
    ("Carga Máxima por Espera", ("Carga Máxima por Espera",), ("Carga Máxima por Espera (kg)", "Carga Máxima por Espera")),
    ("Densidade do Explosivo", ("Densidade do Explosivo",), ("Densidade do Explosivo",)),
    ("Densidade da Rocha", ("Densidade da Rocha",), ("Dens. Rocha (g/cm³)", "Densidade da Rocha", "Dens. Rocha")),
    ("Razão Linear de Carga", ("Razão Linear de Carga Prevista", "Razão Linear de Carga"), ("RL. De Carga (kg/m)", "Razão Linear de Carga")),
    ("Razão de Carga", ("Razão de Carga",), ("RC. Realizada (g/m³)",)),
    ("Volume Total Desmontado", ("Volume Total Desmontado",), ("Volume Total Desm. (m³)", "Volume Total Desmontado")),
    ("Total de Explosivos", ("Total Explosivos Realizado", "Total Explosivos"), ("Total Explosivos Realizados (kg)", "Total Explosivos Realizados")),
)


def clean_norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = "".join(char for char in unicodedata.normalize("NFD", text) if not unicodedata.combining(char))
    text = text.upper()
    return " ".join(re.sub(r"[^A-Z0-9 ]", "", text).split())


def resolve_col(labels: list[str], candidates: tuple[str, ...] | list[str]) -> int:
    normalized = [clean_norm(label) for label in labels]
    for candidate in candidates:
        target = clean_norm(candidate)
        if target in normalized:
            return normalized.index(target)
    for candidate in candidates:
        target = clean_norm(candidate)
        match = next((index for index, label in enumerate(normalized) if target and target in label), -1)
        if match >= 0:
            return match
    return -1


def resolve_material_cols(labels: list[str]) -> list[int]:
    return [
        index
        for index, label in enumerate(labels)
        if clean_norm(label) == "MATERIAL" or re.fullmatch(r"MATERIAL \d+", clean_norm(label))
    ]


def cell_value(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return f"Date({value.year},{value.month - 1},{value.day},{value.hour},{value.minute},{value.second})"
    if isinstance(value, date):
        return f"Date({value.year},{value.month - 1},{value.day})"
    if isinstance(value, time):
        return [value.hour, value.minute, value.second]
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value if math.isfinite(float(value)) else None
    return str(value)


def select_sheet(workbook: Any, preferred: str) -> Any:
    for sheet in workbook.worksheets:
        if sheet.title.casefold() == preferred.casefold():
            return sheet
    for sheet in workbook.worksheets:
        header = [clean_norm(cell.value) for cell in sheet[1]]
        if resolve_col(header, ("Plano",)) >= 0 and resolve_col(header, ("Data",)) >= 0:
            return sheet
    raise ValueError(f"Não encontrei uma aba com as colunas Plano e Data (esperada: {preferred}).")


def build_table(path: Path, side: str, preferred_sheet: str) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"Arquivo de origem não encontrado: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = select_sheet(workbook, preferred_sheet)
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        labels = ["" if value is None else str(value).strip() for value in header_row]
        plan_index = resolve_col(labels, ("Plano",))
        date_index = resolve_col(labels, ("Data",))
        if plan_index < 0 or date_index < 0:
            raise ValueError(f"A aba {sheet.title!r} não possui as colunas Plano e Data.")

        material_indices = resolve_material_cols(labels)
        malha_index = resolve_col(labels, ("Malha Mista", "Malha")) if side == "prev" else -1
        aspect_indices = []
        for _, prev_candidates, real_candidates in ASPECTS:
            candidates = prev_candidates if side == "prev" else real_candidates
            index = resolve_col(labels, candidates)
            if index < 0:
                raise ValueError(f"Coluna necessária não encontrada em {path.name}: {candidates[0]}.")
            aspect_indices.append(index)

        fields: list[tuple[str, int | None, str]] = [
            (labels[plan_index] or "Plano", plan_index, "plain"),
            (labels[date_index] or "Data", date_index, "plain"),
        ]
        if malha_index >= 0:
            fields.append((labels[malha_index] or "Malha Mista", malha_index, "plain"))
        fields.append(("Material", None, "material"))
        for (label, _, _), index in zip(ASPECTS, aspect_indices):
            fields.append((labels[index] or label, index, "plain"))

        columns = [
            {"id": get_column_letter(index + 1), "label": label}
            for index, (label, _, _) in enumerate(fields)
        ]
        rows = []
        valid_dates = []
        for source_row in sheet.iter_rows(min_row=2, values_only=True):
            plan = source_row[plan_index] if plan_index < len(source_row) else None
            when = source_row[date_index] if date_index < len(source_row) else None
            if plan in (None, "") and when in (None, ""):
                continue

            cells = []
            for _, source_index, kind in fields:
                if kind == "material":
                    value = next(
                        (
                            source_row[index]
                            for index in material_indices
                            if index < len(source_row) and source_row[index] not in (None, "")
                        ),
                        None,
                    )
                else:
                    value = source_row[source_index] if source_index is not None and source_index < len(source_row) else None
                value = cell_value(value)
                cells.append(None if value is None else {"v": value})
            rows.append({"c": cells})
            if isinstance(when, (datetime, date)):
                valid_dates.append(when.date() if isinstance(when, datetime) else when)

        if not rows:
            raise ValueError(f"A aba {sheet.title!r} não contém registros de planos.")

        generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
        metadata = {
            "generatedAt": generated_at,
            "sourceFile": path.name,
            "sourceSheet": sheet.title,
            "recordCount": len(rows),
            "firstDate": min(valid_dates).isoformat() if valid_dates else None,
            "lastDate": max(valid_dates).isoformat() if valid_dates else None,
        }
        return {"metadata": metadata, "table": {"cols": columns, "rows": rows}}
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--previsto", required=True, type=Path, help="Caminho para Plano_Fogo_Previsto.xlsx")
    parser.add_argument("--realizado", required=True, type=Path, help="Caminho para Plano_Fogo_Realizado.xlsx")
    parser.add_argument("--output", type=Path, default=Path("data"), help="Pasta de saída dos JSONs")
    args = parser.parse_args()

    previsto = build_table(args.previsto, "prev", "Plano_Fogo_Previsto")
    realizado = build_table(args.realizado, "real", "Desmontes")
    args.output.mkdir(parents=True, exist_ok=True)
    for name, payload in (("previsto", previsto), ("realizado", realizado)):
        target = args.output / f"{name}.json"
        target.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
        print(
            f"{target}: {payload['metadata']['recordCount']} registros, "
            f"{payload['metadata']['firstDate']} a {payload['metadata']['lastDate']}"
        )


if __name__ == "__main__":
    main()
